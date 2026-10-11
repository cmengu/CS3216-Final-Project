"""Read the agreed prices from the Excel file. One price per part."""

import hashlib
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

REQUIRED_HEADERS = ("part_number", "description", "currency", "unit_price")


@dataclass(frozen=True)
class SheetPrice:
    part_number: str
    description: str
    currency: str
    unit_price: Decimal
    unit_price_text: str


@dataclass(frozen=True)
class PriceSheet:
    version: str
    prices: dict[str, SheetPrice]

    def lookup(self, part_number: str | None) -> SheetPrice | None:
        if part_number is None or not str(part_number).strip():
            return None
        return self.prices.get(str(part_number).strip().casefold())


def load_price_sheet(path: Path) -> PriceSheet:
    file = Path(path)
    if not file.is_file():
        raise ValueError(f"Price sheet not found: {file}")
    version = hashlib.sha256(file.read_bytes()).hexdigest()
    workbook = load_workbook(file, data_only=False, read_only=True)
    try:
        if "Prices" not in workbook.sheetnames:
            raise ValueError("Price sheet is missing the Prices worksheet")
        rows = workbook["Prices"].iter_rows()
        try:
            header = next(rows)
        except StopIteration as exc:
            raise ValueError("Price sheet has no header row") from exc
        columns = _columns(header)
        prices: dict[str, SheetPrice] = {}
        for row in rows:
            price = _price(row, columns)
            if price is None:
                continue
            key = price.part_number.strip().casefold()
            prices.setdefault(key, price)
    finally:
        workbook.close()
    return PriceSheet(version=version, prices=prices)


def _columns(header) -> dict[str, int]:
    found = {}
    for index, cell in enumerate(header):
        name = "" if cell.value is None else str(cell.value).strip()
        if name in REQUIRED_HEADERS and name not in found:
            found[name] = index
    missing = [name for name in REQUIRED_HEADERS if name not in found]
    if missing:
        raise ValueError(f"Price sheet is missing the {missing[0]} column")
    return found


def _price(row, columns: dict[str, int]) -> SheetPrice | None:
    part = _text(_cell(row, columns["part_number"]))
    if not part:
        return None
    description = _text(_cell(row, columns["description"]))
    currency = _text(_cell(row, columns["currency"]))
    amount, text = _decimal(_cell(row, columns["unit_price"]))
    return SheetPrice(
        part_number=part,
        description=description,
        currency=currency,
        unit_price=amount,
        unit_price_text=text,
    )


def _cell(row, index: int):
    if index >= len(row):
        return None
    return row[index].value


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _decimal(value) -> tuple[Decimal, str]:
    if isinstance(value, bool) or value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError("Price sheet has a unit price that is not a number")
    if isinstance(value, float):
        text = str(value)
    elif isinstance(value, int) and not isinstance(value, bool):
        text = str(value)
    else:
        text = str(value).strip().replace(",", "")
    try:
        amount = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("Price sheet has a unit price that is not a number") from exc
    if not amount.is_finite():
        raise ValueError("Price sheet has a unit price that is not a number")
    return amount, text
