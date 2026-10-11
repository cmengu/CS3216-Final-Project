"""Generate ten invented invoices and image-only copies from known source records."""

import argparse
import hashlib
import json
from html import escape
from pathlib import Path

from pypdf import PdfReader

from .cases import CATALOG, build_cases

ROOT = Path(__file__).resolve().parents[2]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def invoice_html(case):
    invoice = case["invoice"]
    layout = case["layout"]
    columns = (
        ["part_number", "description", "quantity", "unit_price", "line_total"]
        if layout != 1
        else ["description", "part_number", "unit_price", "quantity", "line_total"]
    )
    labels = {
        "part_number": "Part number",
        "description": "Description",
        "quantity": "Quantity",
        "unit_price": "Unit price",
        "line_total": "Line total",
    }
    rows = "".join(
        "<tr>" + "".join(f'<td class="{key}">{escape(line[key])}</td>' for key in columns) + "</tr>"
        for line in invoice["lines"]
    )
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>
    @page {{ size: A4; margin: 18mm 14mm; @bottom-center {{ content: 'Invented demonstration invoice | Page ' counter(page); font-size: 8pt; color: #555; }} }}
    body {{ font: 10pt {"serif" if layout == 2 else "sans-serif"}; color: #18283b; }}
    h1 {{ font-size: 25pt; margin: 0 0 4mm; }} h2 {{ font-size: 14pt; }}
    .tag {{ font-size: 8pt; letter-spacing: 1pt; color: #555; }}
    header {{ border-bottom: 3px solid {"#193e66" if layout == 0 else "#575348"}; padding-bottom: 5mm; text-align: {"right" if layout == 1 else "left"}; }}
    .details {{ margin: 6mm 0; line-height: 1.7; }}
    table {{ border-collapse: collapse; width: 100%; table-layout: fixed; }}
    thead {{ display: table-header-group; }} tr {{ break-inside: avoid; }}
    th {{ background: #e8edf2; font-size: 9pt; text-align: left; padding: 3mm 1.5mm; }}
    td {{ padding: 3mm 1.5mm; border-bottom: 1px solid #ddd; font-size: 9pt; overflow-wrap: anywhere; }}
    .quantity, .unit_price, .line_total {{ text-align: right; }}
    .totals {{ margin: 7mm 0 0 auto; width: 65mm; break-inside: avoid; }}
    .totals p {{ display: flex; justify-content: space-between; margin: 2mm 0; }}
    .total {{ border-top: 2px solid #18283b; padding-top: 3mm; font-weight: bold; }}
    </style></head><body><header><div class="tag">DUMMY DATA</div>
    <h1>INVOICE</h1><h2>{escape(invoice["supplier_name"])}</h2></header>
    <div class="details">Invoice number: <b>{invoice["invoice_number"]}</b><br>
    Invoice date: {invoice["invoice_date"]}<br>Bill to: Example Semiconductor Company<br>
    Currency: {invoice["currency"]}</div><table><colgroup>
    {"".join(f'<col style="width: {dict(part_number=22, description=31, quantity=14, unit_price=16, line_total=17)[key]}%">' for key in columns)}
    </colgroup><thead><tr>{"".join(f"<th>{labels[key]}</th>" for key in columns)}</tr></thead><tbody>{rows}</tbody></table>
    <div class="totals"><p><span>Subtotal</span><span>{invoice["subtotal"]}</span></p>
    <p><span>Tax (9%)</span><span>{invoice["tax"]}</span></p>
    <p class="total"><span>Total SGD</span><span>{invoice["invoice_total"]}</span></p></div>
    </body></html>"""


def validate_sheet(path):
    # Read-only check of the committed dummy price sheet. Kalent owns the real sheet loader.
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=True, read_only=True)
    try:
        rows = list(book["Prices"].values)
        if rows[0] != ("part_number", "description", "currency", "unit_price"):
            raise ValueError("Unexpected dummy price-sheet columns")
        from decimal import Decimal

        actual = [(r[0], r[1], r[2], Decimal(str(r[3]))) for r in rows[1:]]
        expected = [
            (r["part_number"], r["description"], r["currency"], Decimal(r["unit_price"]))
            for r in CATALOG
        ]
        if actual != expected:
            raise ValueError("Dummy price sheet and source catalogue differ")
    finally:
        book.close()


def generate(directory: Path):
    import pypdfium2 as pdfium
    from PIL import ImageDraw
    from weasyprint import HTML

    directory.mkdir(parents=True, exist_ok=True)
    sheet = directory / "price-sheet.xlsx"
    if not sheet.exists():
        # The authored workbook is a versioned source asset, copied for custom output folders.
        import shutil

        shutil.copyfile(ROOT / "data/dummy/price-sheet.xlsx", sheet)
    validate_sheet(sheet)
    cases = build_cases()
    for case in cases:
        digital = directory / f"{case['id']}.pdf"
        scan = directory / f"{case['id']}-scan.pdf"
        HTML(string=invoice_html(case)).write_pdf(digital, pdf_identifier=case["id"])
        images = []
        with pdfium.PdfDocument(digital) as document:
            for page in document:
                image = page.render(scale=2).to_pil().convert("RGB")
                if case["id"] in ("invoice-06", "invoice-09"):
                    image = image.rotate(0.6, expand=False, fillcolor="white")
                if case["id"] == "invoice-08":
                    ImageDraw.Draw(image).text(
                        (50, image.height - 90), "RECEIVED - DEMO", fill="#777777"
                    )
                images.append(image)
                page.close()
        images[0].save(
            scan,
            "PDF",
            save_all=True,
            append_images=images[1:],
            resolution=144,
            creationDate="D:20261010000000Z",
            modDate="D:20261010000000Z",
        )
        for image in images:
            image.close()
        if any((p.extract_text() or "").strip() for p in PdfReader(scan).pages):
            raise ValueError("Scan fixture must have no embedded text")
        case["files"] = [
            {"name": p.name, "sha256": sha(p), "kind": kind}
            for p, kind in ((digital, "digital"), (scan, "scan"))
        ]
    manifest = {
        "version": "dummy-v1",
        "invented": True,
        "price_sheet_sha256": sha(sheet),
        "catalog": CATALOG,
        "cases": cases,
    }
    (directory / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    return manifest


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, default=ROOT / "data/dummy")
    args = parser.parse_args()
    manifest = generate(args.out)
    print(f"Generated {len(manifest['cases'])} invoices plus scans in {args.out}")


if __name__ == "__main__":
    main()
